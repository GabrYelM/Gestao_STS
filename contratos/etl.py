"""
Módulo de Extração, Transformação e Carga (ETL) dos Contratos de Gestão.

Unifica todas as rotinas de ingestão e normalização de arquivos oficiais:
- Utilitários transparentes de leitura direta de arquivos ou contêineres ZIP
- Catálogo oficial e normalização de Procedimentos (SIGTAP/SMS-SP), CBOs e Estabelecimentos
- Ingestão SIGA (AT-02)
- Ingestão WebSaass (Demonstrativo de Apontamentos Técnicos XML/ZIP)
- Ingestão Visitas Domiciliares Periódicas (eSUS/Centralizador municipal - REL_142)
- Ingestão BI SIGA / SSRS (AT-08, AT-11, AT-39, AT-40, AT-48, AT-49, AT-57, AT-61)
- Ingestão DTIC eMulti (REL_134) e Atendimento Domiciliar eSUS (REL_130)
- Ingestão SISAD Questionário AD (Excel)
"""

import csv
import io
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from contextlib import contextmanager

import openpyxl

# ==============================================================================
# 1. UTILITÁRIOS DE ARQUIVO E ZIP TRANSPARENTE
# ==============================================================================

@contextmanager
def abrir_texto_arquivo_ou_zip(caminho_arquivo, encoding="utf-8-sig", extensao_preferida=".csv"):
    """
    Context manager que entrega um leitor de texto (io.TextIOWrapper ou arquivo normal).
    Se caminho_arquivo for um arquivo .zip ou zip válido:
      - Localiza o primeiro arquivo interno que case com a extensão preferida
        (ex: .csv, .xml, .txt) ignorando pastas do sistema (__MACOSX).
      - Retorna um stream de texto já decodificado no encoding solicitado.
    Caso contrário:
      - Abre o arquivo diretamente do disco com open().
    """
    if isinstance(caminho_arquivo, str) and (
        caminho_arquivo.lower().endswith(".zip") or (os.path.exists(caminho_arquivo) and zipfile.is_zipfile(caminho_arquivo))
    ):
        with zipfile.ZipFile(caminho_arquivo) as z:
            candidatos = [
                n for n in z.namelist()
                if not n.startswith("__MACOSX") and not n.startswith(".") and not n.endswith("/")
            ]
            escolhido = None
            if extensao_preferida:
                for n in candidatos:
                    if n.lower().endswith(extensao_preferida.lower()):
                        escolhido = n
                        break
            if not escolhido and candidatos:
                escolhido = candidatos[0]
            if not escolhido:
                raise ValueError(f"Nenhum arquivo válido encontrado dentro do ZIP: '{caminho_arquivo}'.")

            with z.open(escolhido) as f:
                yield io.TextIOWrapper(f, encoding=encoding, errors="replace")
    else:
        with open(caminho_arquivo, encoding=encoding, errors="replace") as f:
            yield f


@contextmanager
def abrir_binario_arquivo_ou_zip(caminho_arquivo, extensao_preferida=".xlsx"):
    """
    Context manager que entrega um buffer binário ou arquivo aberto em 'rb'.
    Útil para leitura com openpyxl ou ferramentas que exigem stream binário.
    """
    if isinstance(caminho_arquivo, str) and (
        caminho_arquivo.lower().endswith(".zip") or (os.path.exists(caminho_arquivo) and zipfile.is_zipfile(caminho_arquivo))
    ):
        with zipfile.ZipFile(caminho_arquivo) as z:
            candidatos = [
                n for n in z.namelist()
                if not n.startswith("__MACOSX") and not n.startswith(".") and not n.endswith("/")
            ]
            escolhido = None
            if extensao_preferida:
                for n in candidatos:
                    if n.lower().endswith(extensao_preferida.lower()):
                        escolhido = n
                        break
            if not escolhido and candidatos:
                escolhido = candidatos[0]
            if not escolhido:
                raise ValueError(f"Nenhum arquivo válido encontrado dentro do ZIP: '{caminho_arquivo}'.")

            yield io.BytesIO(z.read(escolhido))
    else:
        with open(caminho_arquivo, "rb") as f:
            yield f


# ==============================================================================
# 2. CATÁLOGO E NORMALIZAÇÃO DE PROCEDIMENTOS / CBO / CÓDIGOS
# ==============================================================================

_CATALOGO_PROCEDIMENTOS = None


def carregar_catalogo_procedimentos():
    """Carrega o catálogo geral primário (SIGTAP + procedimentos municipais)."""
    global _CATALOGO_PROCEDIMENTOS
    if _CATALOGO_PROCEDIMENTOS is None:
        caminho = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "services",
            "catalogo_geral.json",
        )
        if os.path.exists(caminho):
            try:
                with open(caminho, encoding="utf-8") as f:
                    dados = json.load(f)
                _CATALOGO_PROCEDIMENTOS = dados.get("procedimentos", {})
            except Exception:
                _CATALOGO_PROCEDIMENTOS = {}
        else:
            _CATALOGO_PROCEDIMENTOS = {}
    return _CATALOGO_PROCEDIMENTOS


def normalizar_cod_procedimento(codigo):
    """
    Normaliza o código de procedimento considerando como primário o catálogo
    extraído do SIGTAP e procedimentos municipais (services/catalogo_geral.json).

    Trata automaticamente:
      - Pontuações e espaços (ex.: '03.01080216' -> '0301080216')
      - Códigos com zeros à esquerda comidos (ex.: '307049086' -> '0307049086')
      - Sufixos/variações do SIGA (ex.: '0301010064A' -> '0301010064')
      - Procedimentos municipais (séries municipais como 0301019..., 0101019..., etc.)
    """
    if not codigo:
        return ""
    texto = str(codigo).strip()
    limpo = re.sub(r"[.\-/\s]", "", texto)

    if len(limpo) == 11 and limpo[:-1].isdigit() and not limpo.isdigit():
        limpo = limpo[:-1]

    if limpo.isdigit() and len(limpo) < 10:
        limpo = limpo.zfill(10)

    catalogo = carregar_catalogo_procedimentos()
    if limpo in catalogo:
        return limpo

    if len(limpo) > 10 and not limpo.isdigit():
        candidato = limpo[:10]
        if candidato in catalogo:
            return candidato

    return limpo


def normalizar_cbo(codigo):
    """Normaliza o código CBO para 6 dígitos."""
    if not codigo:
        return ""
    codigo = str(codigo).strip()
    if codigo.isdigit():
        return codigo.zfill(6)
    return codigo


def normalizar_cnes(codigo):
    """Normaliza o código CNES para 7 dígitos."""
    if not codigo:
        return ""
    codigo = str(codigo).strip()
    if codigo.isdigit():
        return codigo.zfill(7)
    return codigo


def normalizar_cmes(codigo):
    """Normaliza o código CMES para 7 dígitos."""
    if not codigo:
        return ""
    codigo = str(codigo).strip()
    if codigo.isdigit():
        return codigo.zfill(7)
    return codigo


# ==============================================================================
# 3. ETL SIGA AT-02
# ==============================================================================

CABECALHO_ESPERADO_AT02 = "Número_Ano_Mes__AAAAMM_"

MAPA_COLUNAS_AT02 = {
    "Número_Ano_Mes__AAAAMM_": "ano_mes",
    "Código_CNES": "cod_cnes",
    "H1___Nome_Estabelecimento2": "nome_estabelecimento",
    "H1___Código_CMES": "cod_cmes",
    "Tipo_Estabelecimento": "tipo_estabelecimento",
    "Código_CBO_no_SUS": "cod_cbo_sus",
    "Nome_CBO1": "nome_cbo1",
    "Nome_Especialidade2": "nome_especialidade2",
    "Código_Procedimento": "cod_procedimento",
    "Nome_Procedimento2": "nome_procedimento",
    "Nome_Profissional_Siga1": "nome_profissional",
    "Quantidade_Procedimento2": "quantidade",
    "Contagem_Paciente2": "quantidade_pacientes",
}


def _encontrar_cabecalho_at02(linhas):
    for i, linha in enumerate(linhas):
        if linha and linha[0].strip() == CABECALHO_ESPERADO_AT02:
            return i
    raise ValueError(f"Cabeçalho '{CABECALHO_ESPERADO_AT02}' não encontrado no arquivo AT-02.")


def importar_at02(caminho_arquivo, periodo_referencia, db, nome_arquivo=None):
    """
    Lê o CSV do AT-02 e grava em staging_at02.
    Substitui o snapshot anterior desse período por completo.
    """
    with abrir_texto_arquivo_ou_zip(caminho_arquivo, encoding="utf-8-sig", extensao_preferida=".csv") as f:
        linhas = list(csv.reader(f, delimiter=";"))

    idx_cabecalho = _encontrar_cabecalho_at02(linhas)
    cabecalho = [c.strip() for c in linhas[idx_cabecalho]]
    dados = linhas[idx_cabecalho + 1:]

    indices = {}
    for col_csv, col_staging in MAPA_COLUNAS_AT02.items():
        if col_csv in cabecalho:
            indices[col_staging] = cabecalho.index(col_csv)

    fonte_id = db.execute("SELECT id FROM fontes_dados WHERE nome = 'AT02'").fetchone()["id"]

    db.execute("DELETE FROM staging_at02 WHERE ano_mes = ?", (periodo_referencia,))

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    linhas_importadas = 0
    for linha in dados:
        if not linha or len(linha) <= max(indices.values(), default=0):
            continue

        def campo(nome, tipo=str):
            idx = indices.get(nome)
            if idx is None or idx >= len(linha):
                return None
            valor = linha[idx].strip() if linha[idx] else None
            if valor is None or valor == "":
                return None
            return tipo(valor) if tipo is not str else valor

        cod_procedimento = normalizar_cod_procedimento(campo("cod_procedimento"))

        db.execute(
            """INSERT INTO staging_at02 (
                   importacao_id, ano_mes, cod_cnes, nome_estabelecimento, cod_cmes,
                   tipo_estabelecimento, cod_cbo_sus, nome_cbo1, nome_especialidade2, cod_procedimento,
                   nome_procedimento, nome_profissional, quantidade, quantidade_pacientes
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                importacao_id,
                campo("ano_mes"),
                campo("cod_cnes"),
                campo("nome_estabelecimento"),
                campo("cod_cmes"),
                campo("tipo_estabelecimento"),
                campo("cod_cbo_sus"),
                campo("nome_cbo1"),
                campo("nome_especialidade2"),
                cod_procedimento,
                campo("nome_procedimento"),
                campo("nome_profissional"),
                campo("quantidade", int),
                campo("quantidade_pacientes", int),
            ),
        )
        linhas_importadas += 1

    db.execute(
        "UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?",
        (linhas_importadas, importacao_id),
    )
    db.commit()

    return importacao_id, linhas_importadas


# ==============================================================================
# 4. ETL WEBSSAAS (XML/ZIP)
# ==============================================================================

NS_SPREADSHEET = "{urn:schemas-microsoft-com:office:spreadsheet}"

COLUNAS_WEBSSAS = [
    "cod_contrato",
    "contrato",
    "contratada",
    "unidade",
    "periodo",
    "servico",
    "cod_producao",
    "producao",
    "qtde_realizada",
    "qtde_prevista",
]


def _parse_linhas_webssas(xml_path):
    if isinstance(xml_path, str) and (xml_path.lower().endswith(".zip") or (os.path.exists(xml_path) and zipfile.is_zipfile(xml_path))):
        with zipfile.ZipFile(xml_path) as z:
            for name in z.namelist():
                if name.lower().endswith(".xml"):
                    with z.open(name) as xf:
                        tree = ET.parse(xf)
                        break
            else:
                raise ValueError("Nenhum arquivo XML encontrado dentro do ZIP do Webssas.")
    else:
        tree = ET.parse(xml_path)
    root = tree.getroot()

    linhas = []
    for row in root.iter(NS_SPREADSHEET + "Row"):
        row_data = []
        for cell in row.iter(NS_SPREADSHEET + "Cell"):
            index = cell.get(NS_SPREADSHEET + "Index")
            if index:
                index = int(index)
                while len(row_data) < index - 1:
                    row_data.append("")
            data = cell.find(NS_SPREADSHEET + "Data")
            if data is not None:
                if data.get(NS_SPREADSHEET + "Type") == "String" and data.text is not None:
                    row_data.append(" ".join(data.text.split()))
                else:
                    row_data.append(data.text)
            else:
                row_data.append("")
        linhas.append(row_data)
    return linhas


def importar_webssas(caminho_arquivo, periodo_referencia, db, nome_arquivo=None, mes_filtro=None):
    linhas = _parse_linhas_webssas(caminho_arquivo)
    if not linhas:
        raise ValueError("XML do Webssas vazio ou em formato inesperado.")

    dados = linhas[1:]

    fonte_id = db.execute("SELECT id FROM fontes_dados WHERE nome = 'WEBSSAS'").fetchone()["id"]

    from .funcoes import _normalizar_periodo, _periodo_para_webssas
    p_norm = _normalizar_periodo(periodo_referencia)
    p_web = _periodo_para_webssas(periodo_referencia)
    db.execute("DELETE FROM staging_webssas WHERE periodo = ? OR periodo = ?", (p_norm, p_web))

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    linhas_importadas = 0
    for linha in dados:
        valores = dict(zip(COLUNAS_WEBSSAS, linha + [None] * (len(COLUNAS_WEBSSAS) - len(linha))))

        if mes_filtro and (valores.get("periodo") or "").strip() != mes_filtro:
            continue

        def num(campo):
            v = valores.get(campo)
            try:
                return int(v) if v not in (None, "") else None
            except (TypeError, ValueError):
                return None

        db.execute(
            """INSERT INTO staging_webssas (
                   importacao_id, cod_contrato, contrato, contratada, unidade,
                   periodo, servico, cod_producao, producao, qtde_realizada, qtde_prevista
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                importacao_id,
                valores.get("cod_contrato"),
                valores.get("contrato"),
                valores.get("contratada"),
                valores.get("unidade"),
                valores.get("periodo"),
                valores.get("servico"),
                valores.get("cod_producao"),
                valores.get("producao"),
                num("qtde_realizada"),
                num("qtde_prevista"),
            ),
        )
        linhas_importadas += 1

    db.execute(
        "UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?",
        (linhas_importadas, importacao_id),
    )
    db.commit()

    return importacao_id, linhas_importadas


# ==============================================================================
# 5. ETL VISITA DOMICILIAR (eSUS Centralizador municipal)
# ==============================================================================

def importar_visita_domiciliar(caminho_arquivo, periodo_referencia, db, nome_arquivo=None, supervisao_filtro=None):
    with abrir_texto_arquivo_ou_zip(caminho_arquivo, encoding="cp1252", extensao_preferida=".csv") as f:
        linhas = [linha for linha in csv.DictReader(f, delimiter=";") if linha.get("cnes")]

    if not linhas:
        raise ValueError("Arquivo de Visita Domiciliar Periódica vazio ou em formato inesperado.")

    fonte_id = db.execute("SELECT id FROM fontes_dados WHERE nome = 'VISITA_DOMICILIAR'").fetchone()["id"]

    db.execute(
        """DELETE FROM staging_visita_domiciliar WHERE importacao_id IN (
               SELECT id FROM importacoes WHERE fonte_id=? AND periodo_referencia=?)""",
        (fonte_id, periodo_referencia),
    )

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    linhas_importadas = 0
    for linha in linhas:
        if supervisao_filtro and (linha.get("supervisao") or "").strip() != supervisao_filtro.strip():
            continue

        def num(campo):
            v = linha.get(campo)
            try:
                return int(v) if v not in (None, "") else None
            except (TypeError, ValueError):
                return None

        db.execute(
            """INSERT INTO staging_visita_domiciliar (
                   importacao_id, tipo_visita, supervisao, cod_cnes, nome_estabelecimento,
                   cns_profissional, nome_profissional, cod_cbo, nome_cbo, cod_equipe,
                   ano, mes, total_visitas
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                importacao_id,
                linha.get("tipo_visita"),
                linha.get("supervisao"),
                (linha.get("cnes") or "").strip(),
                linha.get("unidade_saude"),
                linha.get("cns_profissional"),
                linha.get("nome_profissional"),
                (linha.get("cod_cbo") or "").strip(),
                linha.get("nome_cbo"),
                linha.get("cod_equipe"),
                linha.get("ano"),
                linha.get("mes"),
                num("total_visitas"),
            ),
        )
        linhas_importadas += 1

    db.execute(
        "UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?",
        (linhas_importadas, importacao_id),
    )
    db.commit()

    return importacao_id, linhas_importadas


# ==============================================================================
# 6. ETL BI SIGA (AT-08, AT-11, AT-39, AT-40, AT-48, AT-49, AT-57, AT-61)
# ==============================================================================

CONFIGS_BI_SIGA = {
    "AT08": {
        "ancora": "H1_Nome_Nivel_21",
        "mapa": {
            "numero_ano_mes": "ano_mes",
            "H1_Nome_Nivel_22": "nivel2",
            "H1_Nome_Nivel_31": "nivel3",
            "H1_Nome_Estabelecimento_Executante1": "nome_estabelecimento",
            "nome_especialidade1": "especialidade",
            "nome_procedimento1": "procedimento_nome",
            "quantidade_procedimento1": "quantidade",
            "quantidade_vaga_ofertada1": "quantidade_vaga_ofertada",
        },
    },
    "AT11": {
        "ancora": "Número_Ano_Mes__AAAAMM_",
        "mapa": {
            "Número_Ano_Mes__AAAAMM_": "ano_mes",
            "H1___Nome_Estabelecimento1": "nome_estabelecimento",
            "Quantidade_Procedimento1": "quantidade",
            "Contagem_Paciente1": "quantidade_pacientes",
        },
    },
    "AT39": {
        "ancora": "H1___Nome_Estabelecimento",
        "mapa": {
            "H1___Nome_Estabelecimento": "nome_estabelecimento",
            "Número_Ano": "ano",
            "Nome_Mes": "mes",
            "Quantidade_Procedimento": "quantidade",
            "Contagem_Paciente": "quantidade_pacientes",
        },
    },
    "AT40": {
        "ancora": "H1___Nome_Estabelecimento",
        "mapa": {
            "H1___Nome_Estabelecimento": "nome_estabelecimento",
            "Nome_CBO": "cbo_nome",
            "Número_Ano": "ano",
            "Nome_Mes": "mes",
            "Quantidade_Procedimento": "quantidade",
            "Contagem_Paciente": "quantidade_pacientes",
        },
    },
    "AT48": {
        "ancora": "Número_Ano_Mes__AAAAMM_",
        "mapa": {
            "Número_Ano_Mes__AAAAMM_": "ano_mes",
            "H1___Nome_Nível_32": "nivel3",
            "H1___Nome_Nível_41": "nivel4",
            "Código_CNES": "cod_cnes",
            "H1___Nome_Estabelecimento1": "nome_estabelecimento",
            "Contagem_Paciente2": "quantidade_pacientes",
        },
    },
    "AT49": {
        "ancora": "H1___Nome_Estabelecimento",
        "mapa": {
            "H1___Nome_Estabelecimento": "nome_estabelecimento",
            "Número_Ano": "ano",
            "Nome_Mes": "mes",
            "Quantidade_Procedimento": "quantidade",
            "Contagem_Paciente": "quantidade_pacientes",
        },
    },
    "AT57": {
        "ancora": "Número_Ano_Mes__AAAAMM_",
        "mapa": {
            "Número_Ano_Mes__AAAAMM_": "ano_mes",
            "H1___Nome_Nível_31": "nivel3",
            "H1___Nome_Nível_21": "nivel2",
            "Código_CNES1": "cod_cnes",
            "H1___Nome_Estabelecimento1": "nome_estabelecimento",
            "Grupo2": "grupo",
            "Código_Procedimento": "procedimento_codigo",
            "Nome_Procedimento2": "procedimento_nome",
            "Quantidade_Procedimento2": "quantidade",
        },
    },
    "AT61": {
        "ancora": "Número_Ano_Mes__AAAAMM_",
        "mapa": {
            "Número_Ano_Mes__AAAAMM_": "ano_mes",
            "H1___Nome_Nível_21": "nivel2",
            "H1___Nome_Nível_31": "nivel3",
            "H1___Nome_Nível_41": "nivel4",
            "Código_CNES1": "cod_cnes",
            "H1___Nome_Estabelecimento1": "nome_estabelecimento",
            "Nome_CBO1": "cbo_nome",
            "Nome_Procedimento1": "procedimento_nome",
            "Quantidade_Procedimento1": "quantidade",
        },
    },
}

ROTULOS_BI_SIGA = {chave: f"SIGA (AT-{chave[2:]})" for chave in CONFIGS_BI_SIGA}


def _parse_numero_bi_siga(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    if not texto:
        return None
    texto = texto.replace(",", "")
    try:
        return float(texto)
    except ValueError:
        return None


def _encontrar_cabecalho_bi_siga(linhas, ancora):
    for i, linha in enumerate(linhas):
        if linha and linha[0].strip() == ancora:
            return i
    raise ValueError(f"Cabeçalho ('{ancora}') não encontrado no arquivo.")


def importar_bi_siga(fonte_at, caminho_arquivo, periodo_referencia, db, nome_arquivo=None):
    if fonte_at not in CONFIGS_BI_SIGA:
        raise ValueError(f"Fonte BI SIGA desconhecida: {fonte_at}")
    config = CONFIGS_BI_SIGA[fonte_at]

    with abrir_texto_arquivo_ou_zip(caminho_arquivo, encoding="utf-8-sig", extensao_preferida=".csv") as f:
        linhas = list(csv.reader(f, delimiter=";"))

    idx_cabecalho = _encontrar_cabecalho_bi_siga(linhas, config["ancora"])
    cabecalho = [c.strip() for c in linhas[idx_cabecalho]]
    dados = linhas[idx_cabecalho + 1:]

    indices = {}
    for col_csv, col_staging in config["mapa"].items():
        if col_csv in cabecalho:
            indices[col_staging] = cabecalho.index(col_csv)

    fonte_id_row = db.execute("SELECT id FROM fontes_dados WHERE nome = ?", (fonte_at,)).fetchone()
    if not fonte_id_row:
        raise ValueError(f"Fonte '{fonte_at}' não cadastrada em fontes_dados.")
    fonte_id = fonte_id_row["id"]

    db.execute(
        "DELETE FROM staging_bi_siga WHERE fonte_at = ? AND ano_mes = ?",
        (fonte_at, periodo_referencia),
    )

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    campos_staging = [
        "ano_mes", "cod_cnes", "nome_estabelecimento", "nivel2", "nivel3", "nivel4",
        "cbo_nome", "especialidade", "procedimento_codigo", "procedimento_nome", "grupo",
        "ano", "mes", "quantidade", "quantidade_pacientes", "quantidade_vaga_ofertada",
    ]

    linhas_importadas = 0
    for linha in dados:
        if not linha or all(not c.strip() for c in linha) or len(linha) <= max(indices.values(), default=0):
            continue

        def campo(nome):
            idx = indices.get(nome)
            if idx is None or idx >= len(linha):
                return None
            valor = linha[idx].strip() if linha[idx] else None
            return valor or None

        valores = {c: campo(c) for c in campos_staging}
        if not valores["ano_mes"]:
            valores["ano_mes"] = periodo_referencia

        for campo_numerico in ("quantidade", "quantidade_pacientes", "quantidade_vaga_ofertada"):
            valores[campo_numerico] = _parse_numero_bi_siga(valores[campo_numerico])

        db.execute(
            f"""INSERT INTO staging_bi_siga (
                    importacao_id, fonte_at, {', '.join(campos_staging)}
                ) VALUES (?, ?, {', '.join(['?'] * len(campos_staging))})""",
            (importacao_id, fonte_at, *[valores[c] for c in campos_staging]),
        )
        linhas_importadas += 1

    db.execute(
        "UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?",
        (linhas_importadas, importacao_id),
    )
    db.commit()

    return importacao_id, linhas_importadas


# ==============================================================================
# 7. ETL OUTRAS FONTES (DTIC REL_134, DTIC REL_130, SISAD)
# ==============================================================================

def _campo_outras(linha, nome):
    valor = linha.get(nome)
    if valor is None:
        return None
    valor = valor.strip() if isinstance(valor, str) else valor
    return valor or None


def importar_dtic_rel134(caminho_arquivo, periodo_referencia, db, nome_arquivo=None):
    with abrir_texto_arquivo_ou_zip(caminho_arquivo, encoding="cp1252", extensao_preferida=".csv") as f:
        leitor = csv.DictReader(f, delimiter=";")
        linhas = list(leitor)

    fonte_id = db.execute("SELECT id FROM fontes_dados WHERE nome = 'DTIC_REL134'").fetchone()["id"]

    db.execute("DELETE FROM staging_dtic_rel134 WHERE periodo_referencia = ?", (periodo_referencia,))

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    total = 0
    for linha in linhas:
        db.execute(
            """INSERT INTO staging_dtic_rel134 (
                   importacao_id, periodo_referencia, coordenadoria, supervisao, oss, tipo_atividade,
                   cnes, unidade, cns_profissional, nome_profissional, cod_cbo, cbo,
                   cod_procedimento, procedimento, data, tema_para_saude, publico_alvo,
                   numero_participantes
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                importacao_id, periodo_referencia,
                _campo_outras(linha, "coordenadoria"), _campo_outras(linha, "supervisao"), _campo_outras(linha, "oss"),
                _campo_outras(linha, "tipo_atividade"), _campo_outras(linha, "cnes"), _campo_outras(linha, "unidade"),
                _campo_outras(linha, "cns_profissional"), _campo_outras(linha, "nome_profissional"),
                _campo_outras(linha, "cod_cbo"), _campo_outras(linha, "cbo"),
                _campo_outras(linha, "cod_procedimento"), _campo_outras(linha, "procedimento"),
                _campo_outras(linha, "data"), _campo_outras(linha, "tema_para_saude"), _campo_outras(linha, "publico_alvo"),
                int(linha["numero_participantes"]) if linha.get("numero_participantes") else None,
            ),
        )
        total += 1

    db.execute("UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?", (total, importacao_id))
    db.commit()
    return importacao_id, total


def importar_dtic_rel130(caminho_arquivo, periodo_referencia, db, nome_arquivo=None):
    with abrir_texto_arquivo_ou_zip(caminho_arquivo, encoding="cp1252", extensao_preferida=".csv") as f:
        leitor = csv.DictReader(f, delimiter=";")
        linhas = list(leitor)

    fonte_id = db.execute("SELECT id FROM fontes_dados WHERE nome = 'DTIC_REL130'").fetchone()["id"]

    db.execute("DELETE FROM staging_dtic_rel130 WHERE periodo_referencia = ?", (periodo_referencia,))

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    total = 0
    for linha in linhas:
        db.execute(
            """INSERT INTO staging_dtic_rel130 (
                   importacao_id, periodo_referencia, coordenadoria, supervisao, oss,
                   cnes, unidade, codigo_atendimento, cns_profissional, nome_profissional,
                   cod_cbo, cbo, nome_equipe, data_cadastro,
                   cod_procedimento, procedimento
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                importacao_id, periodo_referencia,
                _campo_outras(linha, "coordenadoria"), _campo_outras(linha, "supervisao"), _campo_outras(linha, "oss"),
                _campo_outras(linha, "cnes"), _campo_outras(linha, "unidade"), _campo_outras(linha, "codigo_atendimento"),
                _campo_outras(linha, "cns_profissional"), _campo_outras(linha, "nome_profissional"),
                _campo_outras(linha, "cod_cbo"), _campo_outras(linha, "cbo"), _campo_outras(linha, "nome_equipe"),
                _campo_outras(linha, "data_cadastro"),
                _campo_outras(linha, "cod_procedimento"), _campo_outras(linha, "procedimento"),
            ),
        )
        total += 1

    db.execute("UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?", (total, importacao_id))
    db.commit()
    return importacao_id, total


_SISAD_COLUNAS = [
    "ID", "Coordenadoria", "Supervisão", "Unidade", "Ubs de Referência", "Cartão SUS",
    "Nome", "Situação", "Classificação dos pacientes em Cuidados Paliativos",
    "Data da Criação", "Data de Admissao", "Data da Alta", "Motivo da Alta", "Data de Óbito",
]

_SISAD_MAPA_CAMPO = {
    "ID": "id_sisad", "Coordenadoria": "coordenadoria", "Supervisão": "supervisao",
    "Unidade": "unidade", "Ubs de Referência": "ubs_referencia", "Cartão SUS": "cns",
    "Nome": "nome", "Situação": "situacao",
    "Classificação dos pacientes em Cuidados Paliativos": "classificacao_cuidados_paliativos",
    "Data da Criação": "data_criacao", "Data de Admissao": "data_admissao",
    "Data da Alta": "data_alta", "Motivo da Alta": "motivo_alta", "Data de Óbito": "data_obito",
}


def importar_sisad(caminho_arquivo, periodo_referencia, db, nome_arquivo=None):
    with abrir_binario_arquivo_ou_zip(caminho_arquivo, extensao_preferida=".xlsx") as f:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]

        linhas_iter = ws.iter_rows(values_only=True)
        cabecalho = [str(c).strip() if c else "" for c in next(linhas_iter)]
        indices = {col: cabecalho.index(col) for col in _SISAD_COLUNAS if col in cabecalho}

    fonte_id = db.execute("SELECT id FROM fontes_dados WHERE nome = 'SISAD'").fetchone()["id"]

    db.execute("DELETE FROM staging_sisad WHERE periodo_referencia = ?", (periodo_referencia,))

    cursor = db.execute(
        """INSERT INTO importacoes (fonte_id, nome_arquivo, periodo_referencia, status)
           VALUES (?, ?, ?, 'processando')""",
        (fonte_id, nome_arquivo or caminho_arquivo, periodo_referencia),
    )
    importacao_id = cursor.lastrowid

    campos_staging = list(_SISAD_MAPA_CAMPO.values())
    total = 0
    for linha in linhas_iter:
        if linha is None or all(v is None for v in linha):
            continue

        def campo(col):
            idx = indices.get(col)
            if idx is None or idx >= len(linha):
                return None
            valor = linha[idx]
            if isinstance(valor, str):
                valor = valor.strip() or None
            return valor

        valores = {_SISAD_MAPA_CAMPO[col]: campo(col) for col in _SISAD_COLUNAS}

        db.execute(
            f"""INSERT INTO staging_sisad (
                    importacao_id, periodo_referencia, {', '.join(campos_staging)}
                ) VALUES (?, ?, {', '.join(['?'] * len(campos_staging))})""",
            (importacao_id, periodo_referencia, *[valores[c] for c in campos_staging]),
        )
        total += 1

    db.execute("UPDATE importacoes SET status='concluido', linhas_importadas=? WHERE id=?", (total, importacao_id))
    db.commit()
    return importacao_id, total


# ==============================================================================
# ALIASES DE RETROCOMPATIBILIDADE
# ==============================================================================
BI_SIGA_CONFIGS = CONFIGS_BI_SIGA
BI_SIGA_ROTULOS = ROTULOS_BI_SIGA
CONFIGS = CONFIGS_BI_SIGA
ROTULOS = ROTULOS_BI_SIGA

_modulo_atual = sys.modules[__name__]
normalizacao = _modulo_atual
at02 = _modulo_atual
webssas = _modulo_atual
visita_domiciliar = _modulo_atual
bi_siga = _modulo_atual
outras_fontes = _modulo_atual
arquivo_utils = _modulo_atual
